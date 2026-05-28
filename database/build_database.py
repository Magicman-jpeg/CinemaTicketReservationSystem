import json
import sqlite3
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
FRONTEND_DIR = PROJECT_DIR / "frontend"
EXCEL_PATH = Path(r"d:\Users\Owner\Downloads\OOP_GROUP 2_cinema_hall_ticket_sales.xlsx")
DB_PATH = BASE_DIR / "cinema.db"
SCHEMA_PATH = BASE_DIR / "schema.sql"


def normalize(v):
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def dt_to_str(v):
    if v is None:
        return None
    s = str(v)
    return s.replace("T", " ")


def load_sheet_rows(workbook, sheet):
    ws = workbook[sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = [[normalize(c) for c in row] for row in rows[1:] if any(c is not None for c in row)]
    return headers, data


def main():
    if DB_PATH.exists():
        DB_PATH.unlink()

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON;")
    with open(SCHEMA_PATH, "r", encoding="utf-8") as f:
        conn.executescript(f.read())

    tables = ["movie", "screenings", "cinema_seat", "admin", "customer", "transaction"]
    for table in tables:
        headers, rows = load_sheet_rows(wb, table)
        placeholders = ",".join("?" for _ in headers)
        quoted_cols = ",".join(f"[{h}]" for h in headers)
        sql = f"INSERT INTO [{table}] ({quoted_cols}) VALUES ({placeholders})"
        fixed_rows = []
        for r in rows:
            fixed = []
            for col, val in zip(headers, r):
                if "Date" in col or "Time" in col:
                    fixed.append(dt_to_str(val))
                else:
                    fixed.append(val)
            fixed_rows.append(tuple(fixed))
        conn.executemany(sql, fixed_rows)

    conn.execute("""
        INSERT OR IGNORE INTO admin (Admin_ID, Role, Admin_Name, Admin_Username, [Admin _Pass])
        VALUES (9999, 'System Admin', 'Default Admin', 'admin', 'admin123')
    """)
    conn.commit()

    seat_rows = conn.execute("SELECT Seat_No FROM cinema_seat ORDER BY [Row], [Column]").fetchall()
    booked = {
        r[0]
        for r in conn.execute("SELECT DISTINCT Seat_No FROM [transaction] WHERE Payment_Status='Paid'").fetchall()
    }
    seat_json = [{"seatNo": s[0], "booked": s[0] in booked} for s in seat_rows]
    (FRONTEND_DIR / "js" / "seat-data.json").write_text(json.dumps(seat_json, indent=2), encoding="utf-8")

    report = {}
    report["ticketsPerMovie"] = [
        {"movie": r[0], "tickets": r[1]}
        for r in conn.execute("""
            SELECT m.Movie_Title, COUNT(t.Transaction_ID)
            FROM [transaction] t JOIN movie m ON m.Movie_ID = t.Movie_ID
            GROUP BY m.Movie_Title ORDER BY COUNT(*) DESC
        """).fetchall()
    ]
    report["dailyRevenue"] = [
        {"date": r[0], "revenue": r[1]}
        for r in conn.execute("""
            SELECT Transaction_Date, ROUND(SUM(Total_Payment),2)
            FROM [transaction] GROUP BY Transaction_Date ORDER BY Transaction_Date
        """).fetchall()
    ]
    report["weeklyRevenue"] = [
        {"week": r[0], "revenue": r[1]}
        for r in conn.execute("""
            SELECT strftime('%Y-W%W', Transaction_Date), ROUND(SUM(Total_Payment),2)
            FROM [transaction] GROUP BY strftime('%Y-W%W', Transaction_Date)
            ORDER BY strftime('%Y-W%W', Transaction_Date)
        """).fetchall()
    ]
    report["discountUsage"] = [
        {"discountType": r[0], "count": r[1], "totalDiscount": r[2]}
        for r in conn.execute("""
            SELECT COALESCE(Discount_Type,'N/A'), COUNT(*), ROUND(SUM(COALESCE(Discount_Amount,0)),2)
            FROM [transaction] GROUP BY COALESCE(Discount_Type,'N/A')
        """).fetchall()
    ]
    (FRONTEND_DIR / "js" / "report-data.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    conn.close()
    print("Database and frontend JSON generated.")


if __name__ == "__main__":
    main()
